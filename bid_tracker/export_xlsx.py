"""Excel workbook of the whole tracker. openpyxl is imported here only, so the core stays stdlib."""

from __future__ import annotations

import sqlite3
from datetime import date
from pathlib import Path

from bid_tracker.benchmark import load_tabs
from bid_tracker.competitors import head_to_head, profiles
from bid_tracker.db import table_counts
from bid_tracker.rates import load_rates, rate_bands
from bid_tracker.stats import quartiles

SHEETS = ["Dashboard", "Solicitations", "Bids", "Unit Prices", "Rate Cards", "Benchmarks", "Competitors",
          "Ideal Pursuits", "Sources"]
MONEY = '"$"#,##0'
MONEY2 = '"$"#,##0.00'
PCT = "0.0%"
RATIO = "0.00"


def _sheet(wb, title: str, headers: list[str], rows: list[list], formats: dict[int, str] | None = None,
           widths: dict[int, int] | None = None):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for col, fmt in (formats or {}).items():
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for cell in row:
                cell.number_format = fmt
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, max(10, min(40, len(h) + 4)))
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def build(conn: sqlite3.Connection, out: Path, since: str | None = None) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:  # pragma: no cover
        raise SystemExit("export-xlsx needs openpyxl: pip install -r bid_tracker/requirements-dev.txt") from exc

    wb = Workbook()
    wb.remove(wb.active)
    tabs = load_tabs(conn, since=since)

    # Dashboard
    ws = wb.create_sheet("Dashboard")
    ws["A1"] = "FL Public Bid Pricing Tracker"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generated {date.today().isoformat()} from the private bid DB. Every row traces to a public source."
    counts = table_counts(conn)
    r = 4
    for k, v in counts.items():
        ws.cell(r, 1, k.replace("_", " ").title())
        ws.cell(r, 2, v)
        r += 1
    r += 1
    ws.cell(r, 1, "Low-bid tabs by project type (escalated)").font = Font(bold=True)
    r += 1
    for i, h in enumerate(["Project type", "Tabs", "Median low $/SF", "Median low/EE", "Median bidders", "Median gap"], 1):
        ws.cell(r, i, h).font = Font(bold=True)
    r += 1
    for pt in sorted({t["project_type"] for t in tabs}):
        sel = [t for t in tabs if t["project_type"] == pt]
        vals = [pt, len(sel),
                quartiles([t["low_psf_esc"] for t in sel]).get("p50"),
                quartiles([t["low_ee"] for t in sel if t["low_ee"]]).get("p50"),
                quartiles([t["n"] for t in sel]).get("p50"),
                quartiles([t["gap"] for t in sel if t["gap"] is not None]).get("p50")]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.number_format = {3: MONEY2, 4: RATIO, 6: PCT}.get(i, "General")
        r += 1
    ws.column_dimensions["A"].width = 34
    for col in "BCDEF":
        ws.column_dimensions[col].width = 16

    sols = conn.execute("SELECT * FROM v_solicitation_base ORDER BY bid_open_date DESC").fetchall()
    _sheet(wb, "Solicitations",
           ["Ref", "Agency", "Agency type", "County", "Region", "Title", "Method", "Award basis", "Project type",
            "GSF", "Bid open", "Bidders", "Low", "Second", "EE", "EE source", "Award", "Awardee", "Status",
            "Evidence", "Source URL"],
           [[s["sol_ref"], s["agency_name"], s["agency_type"], s["county"], s["region"], s["title"],
             s["procurement_method"], s["award_basis"], s["project_type"], s["gsf"], s["bid_open_date"],
             s["n_bidders"], s["low_bid"], s["second_bid"], s["engineers_estimate"], s["ee_source"],
             s["award_amount"], s["awardee_name"], s["award_status"], s["evidence_class"], s["source_url"]]
            for s in sols],
           {13: MONEY, 14: MONEY, 15: MONEY, 17: MONEY}, {6: 44, 21: 50})

    bids = conn.execute(
        """SELECT s.sol_ref, br.canonical_name, b.bidder_name_raw, b.base_bid, b.total_bid, b.total_basis,
                  b.rank_published, b.responsive, b.is_awardee, br.is_ideal, b.page_ref, b.source_url
           FROM bids b JOIN solicitations s USING (solicitation_id) JOIN bidders br USING (bidder_id)
           ORDER BY s.bid_open_date DESC, s.sol_ref, b.total_bid"""
    ).fetchall()
    _sheet(wb, "Bids", ["Ref", "Bidder", "Name on tab", "Base", "Total", "Basis", "Rank", "Responsive", "Awardee",
                        "Ideal", "Page", "Source URL"],
           [list(b) for b in bids], {4: MONEY, 5: MONEY}, {2: 34, 3: 34, 12: 50})

    items = conn.execute(
        """SELECT s.sol_ref, br.canonical_name, i.line_no, i.item_desc, i.unit, i.qty, i.unit_price, i.extended,
                  i.csi_division, s.bid_open_date, s.region
           FROM bid_items i JOIN bids b USING (bid_id) JOIN solicitations s ON s.solicitation_id = b.solicitation_id
           JOIN bidders br ON br.bidder_id = b.bidder_id ORDER BY s.sol_ref, br.canonical_name, i.line_no"""
    ).fetchall()
    _sheet(wb, "Unit Prices", ["Ref", "Bidder", "Line", "Item", "Unit", "Qty", "Unit price", "Extended", "CSI",
                               "Bid open", "Region"],
           [list(i) for i in items], {7: MONEY2, 8: MONEY}, {4: 44})

    rates = load_rates(conn, since=since)
    _sheet(wb, "Rate Cards", ["Ref", "Bidder", "Service", "Description", "Bucket", "Unit", "Price", "Effective $/SF",
                              "Derived", "After-hours %", "Response hrs", "Awardee", "Ideal", "Source URL"],
           [[x["sol_ref"], x["canonical_name"], x["service"], x["service_desc"], x["bucket"], x["unit"], x["price"],
             x["psf"], "yes" if x["psf_derived"] else "", x["after_hours_premium_pct"], x["response_hrs"],
             x["is_awardee"], x["is_ideal"], x["source_url"]] for x in rates],
           {7: MONEY2, 8: MONEY2}, {14: 50})

    bands = rate_bands(conn, since=since)
    bench_rows = []
    for pt in sorted({(t["project_type"], t["agency_type"], t["region"]) for t in tabs}):
        sel = [t for t in tabs if (t["project_type"], t["agency_type"], t["region"]) == pt]
        q_psf = quartiles([t["low_psf_esc"] for t in sel])
        q_ee = quartiles([t["low_ee"] for t in sel if t["low_ee"]])
        bench_rows.append([*pt, len(sel), q_psf.get("p25"), q_psf.get("p50"), q_psf.get("p75"),
                           q_ee.get("p50"), quartiles([t["n"] for t in sel]).get("p50"),
                           quartiles([t["gap"] for t in sel if t["gap"] is not None]).get("p50"),
                           "SMALL SAMPLE" if len(sel) < 5 else ""])
    for b in bands:
        a = b["all_bidders"]
        bench_rows.append([f"rate:{b['service']}", "", b["bucket"], a["n"], a.get("p25"), a.get("p50"), a.get("p75"),
                           None, None, None, "derived $/SF" if b["derived"] else ""])
    _sheet(wb, "Benchmarks", ["Project type / service", "Agency type", "Region / bucket", "n", "p25 $/SF (esc)",
                              "p50 $/SF (esc)", "p75 $/SF (esc)", "Median low/EE", "Median bidders", "Median gap",
                              "Note"],
           bench_rows, {5: MONEY2, 6: MONEY2, 7: MONEY2, 8: RATIO, 10: PCT}, {1: 24, 3: 18, 11: 16})

    comp = profiles(conn, top=0)
    _sheet(wb, "Competitors", ["Bidder", "Bids", "Wins", "Low bids", "Win rate", "Median % above low",
                               "Median bid/EE", "Rank percentile", "Agencies", "Counties", "Types", "Last bid"],
           [[c["name"], c["bids"], c["wins"], c["low_bids"], c["win_rate"], c["pct_above_low"], c["bid_ee"],
             c["rank_pctile"], ", ".join(c["agencies"]), ", ".join(c["counties"]), ", ".join(c["project_types"]),
             c["last_bid"]] for c in comp],
           {5: PCT, 6: PCT, 7: RATIO, 8: RATIO}, {1: 36, 9: 30})
    h2h = head_to_head(conn)
    if h2h:
        ws = wb["Competitors"]
        start = ws.max_row + 3
        ws.cell(start, 1, "Head to head with Ideal").font = Font(bold=True)
        for i, h in enumerate(["Competitor", "Shared tabs", "Ideal lower", "Median Ideal/them"], 1):
            ws.cell(start + 1, i, h).font = Font(bold=True)
        for j, x in enumerate(h2h, start + 2):
            for i, v in enumerate([x["name"], x["shared_tabs"], x["ideal_lower"], x["ideal_over_them"]], 1):
                ws.cell(j, i, v).number_format = RATIO if i == 4 else "General"

    purs = conn.execute(
        """SELECT s.sol_ref, s.title, p.decision, p.decision_date, p.est_cost_pre_ohp, p.bid_amount, p.markup_pct,
                  p.result, p.rank, p.score_total, p.debrief, p.lessons
           FROM ideal_pursuits p JOIN solicitations s USING (solicitation_id) ORDER BY s.bid_open_date DESC"""
    ).fetchall()
    _sheet(wb, "Ideal Pursuits", ["Ref", "Title", "Decision", "Decided", "Cost pre-OH&P", "Bid", "Markup", "Result",
                                  "Rank", "Score", "Debrief", "Lessons"],
           [list(p) for p in purs], {5: MONEY, 6: MONEY, 7: PCT}, {2: 40, 11: 50, 12: 50})

    srcs = conn.execute(
        """SELECT s.sol_ref, x.evidence_class, x.retrieved_at, x.url, x.excerpt
           FROM solicitation_sources x JOIN solicitations s USING (solicitation_id) ORDER BY s.sol_ref"""
    ).fetchall()
    _sheet(wb, "Sources", ["Ref", "Evidence", "Retrieved", "URL", "Excerpt"], [list(x) for x in srcs],
           widths={4: 60, 5: 80})

    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
