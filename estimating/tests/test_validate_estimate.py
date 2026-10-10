"""Tests for the provenance-aware lineitems.csv schema in validate_estimate.py and
build_estimate_xlsx.py (12-column core, optionally + the 5-column provenance tail)."""

import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
VALIDATE = ROOT / "estimating" / "scripts" / "validate_estimate.py"
BUILD = ROOT / "estimating" / "scripts" / "build_estimate_xlsx.py"

CORE = "division,section,item,description,qty,unit,unit_mat,unit_lab,unit_equip,unit_sub,waste_pct,notes"
TAIL = "line_id,source_sheet,method,confidence,price_basis"
MARKUPS = """key,value
material_sales_tax_pct,7
general_conditions_pct,10
contingency_pct,3
insurance_pct,1
bond_pct,0
permit_pct,1
ohp_pct,8
"""


def write(proj, header, rows):
    proj.mkdir(parents=True, exist_ok=True)
    (proj / "lineitems.csv").write_text("\n".join([header, *rows]) + "\n", encoding="utf-8")
    (proj / "markups.csv").write_text(MARKUPS, encoding="utf-8")


def run(script, proj):
    p = subprocess.run([sys.executable, str(script), str(proj)], capture_output=True, text=True)
    return p.returncode, p.stdout + p.stderr


def test_core_header_still_valid(tmp_path):
    write(tmp_path, CORE, ["03,03 30 00,SOG,5in SOG,4500,SF,2.85,1.10,0.20,,7,S-201"])
    code, out = run(VALIDATE, tmp_path)
    assert code == 0, out
    assert "12 columns, header exact" in out
    assert "no provenance columns" in out


def test_provenance_header_valid(tmp_path):
    write(tmp_path, f"{CORE},{TAIL}",
          ["03,03 30 00,SOG,5in SOG,4500,SF,2.85,1.10,0.20,,7,S-201,03-001,S-201,measured,med-high,budgetary"])
    code, out = run(VALIDATE, tmp_path)
    assert code == 0, out
    assert "17 columns, header exact (with provenance)" in out
    assert "[PASS] price-basis" in out


def test_priced_line_without_basis_warns(tmp_path):
    write(tmp_path, f"{CORE},{TAIL}",
          ["03,03 30 00,SOG,5in SOG,4500,SF,2.85,1.10,0.20,,7,S-201,03-001,S-201,measured,med-high,"])
    code, out = run(VALIDATE, tmp_path)
    assert "[WARN] price-basis" in out


def test_bad_provenance_values_warn(tmp_path):
    write(tmp_path, f"{CORE},{TAIL}",
          ["03,03 30 00,SOG,5in SOG,4500,SF,2.85,1.10,0.20,,7,S-201,03-001,S-201,eyeballed,certain,guess"])
    code, out = run(VALIDATE, tmp_path)
    assert "[WARN] provenance" in out


def test_partial_tail_is_schema_fail(tmp_path):
    write(tmp_path, f"{CORE},line_id",
          ["03,03 30 00,SOG,5in SOG,4500,SF,2.85,1.10,0.20,,7,S-201,03-001"])
    code, out = run(VALIDATE, tmp_path)
    assert code == 1
    assert "header mismatch" in out


def test_xlsx_build_keeps_provenance_and_ties_out(tmp_path):
    pytest.importorskip("openpyxl")
    write(tmp_path, f"{CORE},{TAIL}", [
        "03,03 30 00,SOG,5in SOG,4500,SF,2.85,1.10,0.20,,7,S-201,03-001,S-201,measured,med-high,budgetary",
        "08,08 51 13,Impact windows,NOA windows,38,EA,,,,1450,0,sub,08-001,A-501,counted,med-high,quote",
    ])
    code, out = run(BUILD, tmp_path)
    assert code == 0, out
    from openpyxl import load_workbook
    det = load_workbook(tmp_path / "estimate.xlsx")["Detail"]
    assert [det.cell(1, c).value for c in range(17, 23)] == [
        "Notes", "Line ID", "Source Sheet", "Method", "Confidence", "Price Basis"]
    code, out = run(VALIDATE, tmp_path)
    assert "[PASS] xlsx-tie-out" in out, out
